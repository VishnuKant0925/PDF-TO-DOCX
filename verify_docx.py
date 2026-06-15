"""
verify_docx.py — Verify DOCX output by printing entry details.

Quick verification script that reads the generated DOCX and prints
the content structure so we can confirm:
    - Two-column layout is set
    - Entries are present with all three fields
    - Hindi text is properly stored in Unicode
    - Font settings are correct

Usage:
    venv\\Scripts\\python.exe verify_docx.py test_page1_v2.docx
"""

import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')

from docx import Document
from docx.oxml.ns import qn


def verify_docx(path: str):
    doc = Document(path)

    print(f"\n{'='*62}")
    print(f"  DOCX Verification: {path}")
    print(f"{'='*62}\n")

    # Check sections
    print(f"  Sections: {len(doc.sections)}")
    for i, section in enumerate(doc.sections):
        print(f"\n  --- Section {i+1} ---")
        print(f"  Page: {section.page_width.cm:.1f} x {section.page_height.cm:.1f} cm")
        print(f"  Margins: L={section.left_margin.cm:.1f} R={section.right_margin.cm:.1f} "
              f"T={section.top_margin.cm:.1f} B={section.bottom_margin.cm:.1f} cm")

        # Check columns
        sect_pr = section._sectPr
        cols_elem = sect_pr.find(qn('w:cols'))
        if cols_elem is not None:
            num_cols = cols_elem.get(qn('w:num'), '1')
            spacing = cols_elem.get(qn('w:space'), 'N/A')
            print(f"  Columns: {num_cols} (spacing: {spacing})")
        else:
            print(f"  Columns: 1 (default)")

    # Check paragraphs
    print(f"\n  Total paragraphs: {len(doc.paragraphs)}")

    # Count entries (paragraphs with bold runs = dictionary entries)
    entries_count = 0
    hindi_count = 0
    english_count = 0
    column_breaks = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            # Check for column break
            for run in para.runs:
                for br in run._element.findall(qn('w:br')):
                    if br.get(qn('w:type')) == 'column':
                        column_breaks += 1
            continue

        has_bold = any(run.bold for run in para.runs if run.text.strip())
        has_hindi = any('\u0900' <= c <= '\u097F' for c in text)

        if has_bold and len(text) > 2:
            entries_count += 1
        if has_hindi:
            hindi_count += 1

        # Count English text
        if any(c.isascii() and c.isalpha() for c in text):
            english_count += 1

    print(f"  Dictionary entries (bold): {entries_count}")
    print(f"  Paragraphs with Hindi: {hindi_count}")
    print(f"  Paragraphs with English: {english_count}")
    print(f"  Column breaks: {column_breaks}")

    # Print first 10 entries for spot-check
    print(f"\n  --- First 10 entries ---")
    entry_num = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        has_bold = any(run.bold for run in para.runs if run.text.strip())
        if has_bold and len(text) > 2:
            entry_num += 1
            if entry_num <= 10:
                # Show run details
                runs_info = []
                for run in para.runs:
                    if run.text.strip():
                        font_name = run.font.name or "?"
                        bold = "B" if run.bold else ""
                        italic = "I" if run.italic else ""
                        style = f"[{font_name} {bold}{italic}]".strip()
                        runs_info.append(f'"{run.text}" {style}')
                print(f"  {entry_num:2d}. {' | '.join(runs_info)}")

    print(f"\n  Total entries: {entry_num}")
    print()


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "test_page1_v2.docx"
    verify_docx(path)
